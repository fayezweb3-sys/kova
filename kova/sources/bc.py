"""BC PharmaCare: weekly downloadable data file, plus the monthly LCA and RDP files.

pddf.zip holds one CSV with a header row: one record per DIN, per plan, per period,
with two years of history (layout: "PharmaCare Downloadable Data File, Interpretation
Document"). A record is current when today falls between Rec Eff Date and Rec End Date.

The LCA and RDP master spreadsheets (Excel) are matched by header name, so a moved
or added column doesn't break the load (layout: "LCA/RDP Data Dictionary").
"""

from __future__ import annotations

import csv
import io
import re
import sqlite3
import zipfile
from datetime import date
from pathlib import Path

from .. import db as D
from .. import normalize as N

PDDF_URL = "https://www.health.gov.bc.ca/pharmacare/outgoing/pddf.zip"
LCA_PAGE_URL = ("https://www2.gov.bc.ca/gov/content/health/practitioner-professional-resources/pharmacare/"
                "pharmacies/low-cost-alternative-lca-and-reference-drug-program-rdp-data-files")
LCA_LINK = re.compile(r'href="([^"]*lca_current[^"]*\.xlsx?)"', re.I)
RDP_LINK = re.compile(r'href="([^"]*rdp_current[^"]*\.xlsx?)"', re.I)


def _key(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")


# pddf header name (normalised) -> bc_coverage column
PDDF_MAP = {
    "din_pin": "din", "plan": "plan", "rec_eff_date": "eff_date", "rec_end_date": "end_date",
    "ben_grp_list": "benefit_groups", "lca_ind": "lca_ind", "pay_gen_ind": "pay_generic",
    "brand_nm": "brand_name", "manuf": "manufacturer", "generic_nm": "generic_name",
    "dosage_form": "dosage_form", "max_price": "max_price", "lca_price": "lca_price",
    "rdp_cat": "rdp_cat", "rdp_sub_cat": "rdp_subcat", "rdp_price": "rdp_price",
    "can_fed_reg_cd": "fed_reg_code", "pcare_plan_desc": "plan_desc", "max_days_supply": "max_days_supply",
    "qty_limit": "qty_limit", "formulary_list_date": "list_date", "ltd_use_flag": "limited_use",
}
COV_COLS = ["din", "plan", "plan_desc", "eff_date", "end_date", "benefit_groups", "lca_ind", "pay_generic",
            "brand_name", "manufacturer", "generic_name", "dosage_form", "fed_reg_code", "max_price", "lca_price",
            "rdp_cat", "rdp_subcat", "rdp_price", "max_days_supply", "qty_limit", "list_date", "limited_use",
            "is_current"]
_NUM = {"max_price", "lca_price", "rdp_price"}
_INT = {"max_days_supply", "qty_limit"}
_DATES = {"eff_date", "end_date", "list_date"}


def _read_pddf(path: Path) -> str:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist() if n.lower().endswith((".csv", ".txt"))]
            if not names:
                raise FileNotFoundError(f"No CSV inside {path}")
            raw = zf.read(max(names, key=lambda n: zf.getinfo(n).file_size))
    else:
        raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def parse_pddf(path: Path, on: date | None = None) -> list[dict]:
    today = (on or date.today()).isoformat()
    reader = csv.reader(io.StringIO(_read_pddf(path)))
    header = [_key(h) for h in next(reader)]
    idx = {PDDF_MAP[h]: i for i, h in enumerate(header) if h in PDDF_MAP}
    if "din" not in idx or "plan" not in idx:
        raise ValueError(f"{path}: expected 'DIN/PIN' and 'Plan' columns, found {header[:8]}")
    rows = []
    for raw in reader:
        if not raw or not any(c.strip() for c in raw):
            continue
        r = {}
        for col, i in idx.items():
            v = raw[i].strip() if i < len(raw) else ""
            if col in _NUM:
                v = N.money(v)
            elif col in _INT:
                v = int(N.money(v)) if N.money(v) is not None else None
            elif col in _DATES:
                v = N.yyyymmdd(v)
            elif col == "din":
                v = N.din(v)
            r[col] = v if v != "" else None
        if not r.get("din"):
            continue
        eff, end = r.get("eff_date") or "0000-00-00", r.get("end_date") or "9999-12-31"
        r["is_current"] = int(eff <= today <= end)
        rows.append(r)
    return rows


def _sheet_rows(path: Path) -> list[dict]:
    """Rows from the first sheet (xlsx) or a CSV, keyed by normalised header. Finds the header row by 'DIN'."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        grid = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        grid = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig", errors="replace"))))
    for i, row in enumerate(grid[:30]):
        keys = [_key(c) for c in row]
        if "din" in keys:
            return [dict(zip(keys, (c.strip() for c in r))) for r in grid[i + 1:] if any(str(c).strip() for c in r)]
    raise ValueError(f"{path}: no header row with a DIN column")


def _pick(row: dict, *names: str) -> str | None:
    for n in names:
        if row.get(n):
            return row[n]
    return None


def parse_lca(path: Path) -> list[dict]:
    out = []
    for r in _sheet_rows(path):
        din = N.din(r.get("din"))
        if not din:
            continue
        out.append(dict(
            din=din, lca_category=_pick(r, "lca_category_no", "lca_category", "lca_cat", "category"),
            chemical_name=_pick(r, "chemical_name", "generic_name"), drug_name=_pick(r, "drug_name"),
            manufacturer=_pick(r, "man", "manufacturer"),
            full_partial=_pick(r, "full_partial_lca_benefit", "full_partial"),
            plan_b_only=_pick(r, "plan_b_only"), rdp=_pick(r, "rdp"), special_auth=_pick(r, "sa"),
            provisional=_pick(r, "prov"), max_price=N.money(_pick(r, "max_price")), lca_price=N.money(_pick(r, "lca_price")),
        ))
    return out


def parse_rdp(path: Path) -> list[dict]:
    out = []
    for r in _sheet_rows(path):
        din = N.din(r.get("din"))
        if din:
            out.append(dict(din=din, rdp_category=_pick(r, "rdp_category", "rdp_cat"), generic_name=_pick(r, "generic_name"),
                            drug_name=_pick(r, "drug_name"), manufacturer=_pick(r, "man"), rdp=_pick(r, "rdp")))
    return out


def _insert(conn, table, rows, cols):
    conn.execute(f"DELETE FROM {table}")
    conn.executemany(f"INSERT INTO {table} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                     [[r.get(c) for c in cols] for r in rows])


def load(conn: sqlite3.Connection, pddf: Path, lca: Path | None = None, rdp: Path | None = None) -> dict:
    cov = parse_pddf(pddf)
    if not cov:
        raise ValueError(f"No PharmaCare records found in {pddf}")
    _insert(conn, "bc_coverage", cov, COV_COLS)
    out = {"records": len(cov), "current": sum(r["is_current"] for r in cov)}
    if lca:
        rows = parse_lca(lca)
        _insert(conn, "bc_lca", rows, ["din", "lca_category", "chemical_name", "drug_name", "manufacturer", "full_partial",
                                      "plan_b_only", "rdp", "special_auth", "provisional", "max_price", "lca_price"])
        out["lca"] = len(rows)
    if rdp:
        rows = parse_rdp(rdp)
        _insert(conn, "bc_rdp", rows, ["din", "rdp_category", "generic_name", "drug_name", "manufacturer", "rdp"])
        out["rdp"] = len(rows)
    D.log_run(conn, "bc", str(pddf), out["current"], max((r["eff_date"] or "" for r in cov), default=None))
    return out
